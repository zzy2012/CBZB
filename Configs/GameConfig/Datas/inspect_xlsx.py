from openpyxl import load_workbook
for f in ["item.xlsx","__tables__.xlsx","__beans__.xlsx"]:
    wb = load_workbook(f)
    print("=====", f, wb.sheetnames)
    for ws in wb.worksheets:
        print(" -- sheet", ws.title, ws.max_row, "x", ws.max_column)
        for i,row in enumerate(ws.iter_rows(min_row=1, max_row=min(6,ws.max_row), values_only=True)):
            print("   ", i+1, row)
